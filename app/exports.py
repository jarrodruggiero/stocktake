"""CSV / Excel exports.

  transactions      every buy/sell/DRP — the raw ledger
  holdings          current valuation per open position
  dividends         income by payment, with franking credits and the gross-up
  realised_cgt      disposals matched FIFO to their parcels
  unrealised_cgt    what a sale today would look like, per parcel still held
  performance       per-FY invested / value / gain
  closed_positions  one row per instrument fully exited

Every money column is AUD at each row's own FX rate unless the header says
"native". The discount columns apply the 50% CGT discount per parcel, before
losses are netted — the FY report nets losses first and is the number to lodge,
and that caveat travels with the file (decisions.md #77).
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from . import clock, cpi, fyreport, queries
from .models import CorporateAction, Dividend, Instrument, Trade

ZERO = Decimal(0)
Report = tuple[list[str], list[list]]

# report key -> (title, builder). `fy` is passed to the ones that take a year.
TITLES = {
    "transactions": "All trades",
    "holdings": "Current holdings",
    "dividends": "Dividends and franking",
    "realised_cgt": "Realised capital gains",
    "unrealised_cgt": "Unrealised capital gains",
    "performance": "Performance by financial year",
    "closed_positions": "Closed positions",
}


def _d(value: Decimal | float | None, places: str = "0.01"):
    if value is None:
        return ""
    return Decimal(value).quantize(Decimal(places))


def _total(parts) -> Decimal | None:
    """The sum, or None when any part is unknown: a total short of a part it
    cannot convert is a wrong number, and a blank is an honest one (#5)."""
    parts = list(parts)
    return None if any(p is None for p in parts) else sum(parts, ZERO)


def _aud(value: Decimal, fx: Decimal | None) -> Decimal | None:
    """Convert, or return None so the cell comes out blank.

    Multiplying by 1 when the rate is unknown would put a native-currency
    number under a column headed "(AUD)", and a spreadsheet will happily total
    that column. Blank is the honest answer; the caller resolves what it can
    through `queries.FxBook` before getting here.
    """
    if fx is None:
        return None
    return value * fx


# --------------------------------------------------------------------------- #
# Reports
# --------------------------------------------------------------------------- #

def transactions(session: Session, ticker: str | None = None, **_) -> Report:
    headers = [
        "Date", "Ticker", "Name", "Exchange", "Type", "Units",
        "Price (native)", "Brokerage (native)", "Currency", "FX rate to AUD",
        "Value (AUD)", "Recorded by", "Note",
    ]
    stmt = select(Trade).options(selectinload(Trade.instrument)).order_by(Trade.date, Trade.id)
    fxbook = queries.FxBook(session)
    rows = []
    for t in session.scalars(stmt):
        if ticker and t.instrument.ticker != ticker:
            continue
        gross = t.quantity * t.unit_price
        signed = gross + t.brokerage if t.type in ("buy", "drp") else gross - t.brokerage
        # The FX column shows what the trade actually captured; the AUD column
        # may additionally fall back to the stored series for that date.
        rows.append(((t.date, 1, t.id), [
            t.date.isoformat(), t.instrument.ticker, t.instrument.name or "",
            t.instrument.exchange, t.type, _d(t.quantity, "0.00000001"),
            _d(t.unit_price, "0.000001"), _d(t.brokerage), t.instrument.currency,
            _d(t.fx_rate, "0.000001"),
            _d(_aud(signed, fxbook.of(t, t.instrument.currency))),
            t.user_id or "", t.note or "",
        ]))
    # Corporate actions in the same ledger, each before its day's trades (#146).
    for a in session.scalars(select(CorporateAction)
                             .options(selectinload(CorporateAction.instrument))):
        if ticker and a.instrument.ticker != ticker:
            continue
        money = None if a.amount is None else _aud(a.amount, fxbook.of(a, a.instrument.currency))
        rows.append(((a.date, 0, a.id), [
            a.date.isoformat(), a.instrument.ticker, a.instrument.name or "",
            a.instrument.exchange, queries.action_label(a), "", "", "",
            a.instrument.currency, _d(a.fx_rate, "0.000001"),
            "" if a.amount is None else _d(money), a.user_id or "", a.note or "",
        ]))
    return headers, [row for _order, row in sorted(rows, key=lambda pair: pair[0])]


def holdings(session: Session, **_) -> Report:
    headers = [
        "Ticker", "Name", "Class", "Currency", "Units", "Avg price (native)",
        "Last price (native)", "Price date", "Cost base (AUD)", "Value (AUD)",
        "Unrealised gain (AUD)", "Gain %", "Dividends (native)", "DRP", "Note",
    ]
    open_positions, _closed = queries.split_positions(queries.all_holdings(session))
    rows = []
    for h in open_positions:
        rows.append([
            h.instrument.ticker, h.instrument.name or "", h.instrument.asset_class,
            h.instrument.currency, _d(h.units, "0.00000001"),
            _d(h.avg_price, "0.000001"), _d(h.price, "0.000001"),
            h.price_date.isoformat() if h.price_date else "",
            _d(h.cost_aud), _d(h.value_aud), _d(h.gain_aud),
            _d(h.gain_pct * 100, "0.01") if h.gain_pct is not None else "",
            _d(h.dividends_cash), "yes" if h.drp else "no", h.note or "",
        ])
    return headers, rows


def dividends(session: Session, fy: int | None = None, **_) -> Report:
    headers = [
        "Payment date", "Period ended", "Ticker", "Name", "Cash (native)", "Currency",
        "FX rate to AUD", "Cash (AUD)", "Franking credits (AUD)",
        "Foreign tax withheld (native)", "Grossed up (AUD)", "Reinvested",
        "Financial year", "Note",
    ]
    stmt = (
        select(Dividend)
        .options(selectinload(Dividend.instrument))
        .order_by(Dividend.date, Dividend.id)
    )
    fxbook = queries.FxBook(session)
    rows = []
    for d in session.scalars(stmt):
        year = fyreport.current_fy(fyreport.income_date(d))
        if fy and year != fy:
            continue
        rate = fxbook.of(d, d.instrument.currency)
        cash_aud = _aud(d.cash_amount, rate)
        withheld_aud = _aud(d.withholding_tax or ZERO, rate)
        # NOT `or ZERO`. `franking_credits` is None when no statement has ever
        # supplied one and 0 when the distribution was unfranked, and ABOUT
        # (below) promises the reader those are different cells. Flattening
        # them here turned "we don't know" into "there was none" — a tax figure
        # somebody would have taken at face value.
        franking = d.franking_credits
        rows.append([
            d.date.isoformat(), d.period_end.isoformat() if d.period_end else "",
            d.instrument.ticker, d.instrument.name or "",
            _d(d.cash_amount, "0.0001"), d.instrument.currency,
            _d(d.fx_rate, "0.000001"), _d(cash_aud), _d(franking),
            _d(d.withholding_tax, "0.0001"),
            # No AUD cash figure means no honest total — and neither does an
            # unrecorded credit, for the same reason: the grossed-up amount IS
            # cash + franking + tax withheld, so it cannot be known while part
            # of it is not.
            _d(None if cash_aud is None or franking is None
               else cash_aud + franking + withheld_aud),
            "yes" if d.reinvest_trade_id else "no",
            fyreport.fy_label(year), d.note or "",
        ])
    return headers, rows


# What made a row of the realised gains: a sale, or a capital change beyond a
# parcel's cost base, which sold nothing (decisions.md #146).
EVENTS = {"sale": "sale", "capital_return": "return of capital",
          "cost_base": "cost base decrease"}


def realised_cgt(session: Session, fy: int | None = None, **_) -> Report:
    """One row per parcel consumed by a sale — the tax-time report.

    "Gain after discount" is the per-parcel figure: half the gain when the
    parcel was held more than twelve months, and from 1 July 2027 half the gain
    to 30 June 2027 only (decisions.md #150). The FY report applies losses
    before the discount, so use its total when lodging.
    """
    headers = [
        "Sale date", "Ticker", "Name", "Units", "Acquired", "Days held",
        "Cost base (AUD)", "Proceeds (AUD)", "Gross gain (AUD)",
        "Discount eligible", "CGT discount (AUD)", "Gain after discount (AUD)",
        "Financial year", "Event", "Value at 30 Jun 2027 (AUD)", "Gain to 30 Jun 2027 (AUD)",
        "Indexation factor", "Indexed cost base (AUD)", "Indexation (AUD)",
    ]
    instruments = (
        session.scalars(select(Instrument).options(selectinload(Instrument.trades),
                                                   selectinload(Instrument.actions)))
        .unique()
        .all()
    )
    fxbook = queries.FxBook(session)
    rows = []
    for inst in instruments:
        # A sale, or a capital change that can go beyond a cost base (#146).
        if not (any(t.type == "sell" for t in inst.trades)
                or any(a.kind != "split" for a in inst.actions)):
            continue
        for disposal in fyreport.instrument_disposals(inst, fxbook):
            year = disposal.date.year + (0 if disposal.date.month <= 6 else 1)
            if fy and year != fy:
                continue
            for parcel in disposal.parcels:
                gain = parcel.gain          # None when either side had no rate
                discount = parcel.discount
                rows.append([
                    disposal.date.isoformat(), inst.ticker, inst.name or "",
                    _d(parcel.quantity, "0.00000001"), parcel.acquired.isoformat(),
                    (disposal.date - parcel.acquired).days,
                    _d(parcel.cost), _d(parcel.proceeds), _d(gain),
                    "yes" if parcel.discountable else "no",
                    _d(discount), _d(None if gain is None else gain - discount),
                    fyreport.fy_label(year), EVENTS[disposal.event],
                    _d(parcel.deferred_proceeds), _d(parcel.deferred),
                    _d(parcel.factor, "0.001"),
                    _d(parcel.indexed_cost if parcel.factor is not None else None),
                    _d(parcel.indexation),
                ])
    rows.sort(key=lambda r: (r[0], r[1]))
    return headers, rows


def unrealised_cgt(session: Session, **_) -> Report:
    """What a sale at today's price would realise, per parcel still held:
    worked as the sale simulator works one, so from 1 July 2027 the latest CPI
    stands in for the quarter's, which is not out yet (decisions.md #150)."""
    headers = [
        "Ticker", "Name", "Acquired", "Units", "Days held", "Cost base (AUD)",
        "Market value (AUD)", "Unrealised gain (AUD)", "Discount eligible",
        "Gain after discount (AUD)", "Indexation (AUD)",
    ]
    today = clock.today()
    instruments = (
        session.scalars(select(Instrument).options(selectinload(Instrument.trades),
                                                   selectinload(Instrument.actions)))
        .unique()
        .all()
    )
    fxbook = queries.FxBook(session)
    fxbook.cpi = cpi.CpiBook(session, estimate=True)
    rows = []
    for inst in instruments:
        latest = queries._latest_price(session, inst.id)
        if latest is None:
            continue
        price, _date = latest
        fx_now = fxbook.rate(inst.currency, today)
        # The same FIFO walk as the CGT engine and the holdings, so a holding's
        # unrealised gain is the sum of these rows (decisions.md #54).
        for parcel in fyreport.walk(inst, fxbook).held:
            units = parcel.remaining
            # A spreadsheet will happily add a USD number to an AUD column, so
            # a figure we can't convert is left EMPTY rather than converted at
            # 1:1 — the same refusal the dashboard makes via `totals.excluded`.
            cost = parcel.cost_held
            value = None if fx_now is None else units * price * fx_now
            # `cost` is None only for a currency with no rate stored at all,
            # and then there is no rate for `value` either.
            gain = None if value is None or cost is None else value - cost
            sold = fyreport.as_if_sold(parcel, value, today, fxbook)
            rows.append([
                inst.ticker, inst.name or "", parcel.acquired.isoformat(),
                _d(units, "0.00000001"), (today - parcel.acquired).days,
                _d(cost), _d(value), _d(gain),
                "yes" if sold.discountable else "no",
                "" if sold.gain is None else _d(sold.gain - sold.discount),
                _d(sold.indexation),
            ])
    rows.sort(key=lambda r: (r[0], r[2]))
    return headers, rows


def performance(session: Session, **_) -> Report:
    headers = [
        "Financial year", "Invested in year (AUD)", "Invested cumulative (AUD)",
        "Value at year end (AUD)", "Total return cumulative (AUD)", "Return %",
        "Return in year (AUD)", "Year return %", "Note",
    ]
    annual = queries.annual_series(session)
    rows = []
    for i, label in enumerate(annual["labels"]):
        left_out = ", ".join(f"{ticker} ({why})" for ticker, why in annual["left_out"][i])
        rows.append([
            label, _d(annual["invested_yr"][i]), _d(annual["invested_cum"][i]),
            _d(annual["value_eofy"][i]), _d(annual["gain_cum"][i]),
            _d(annual["gain_pct"][i] * 100), _d(annual["gain_yr"][i]),
            _d(annual["yr_pct"][i] * 100),
            f"Leaves out {left_out}" if left_out else "",
        ])
    return headers, rows


def closed_positions(session: Session, **_) -> Report:
    """Fully exited instruments — outlay, proceeds and realised profit.

    Capital-gain columns come from the FIFO engine, so they line up with the
    realised CGT report rather than being a second opinion.
    """
    headers = [
        "Ticker", "Name", "Units sold", "First bought", "Last sold",
        "Outlay incl. brokerage (AUD)", "Proceeds net of brokerage (AUD)",
        "Dividends (AUD)", "Gross capital gain (AUD)", "CGT discount (AUD)",
        "Capital gain after discount (AUD)", "Total return incl. dividends (AUD)",
        "Note",
    ]
    _open, closed = queries.split_positions(queries.all_holdings(session))
    prefs = queries.prefs_by_instrument(session)
    fxbook = queries.FxBook(session)
    rows = []
    for position in closed:
        inst = position.instrument
        disposals = fyreport.instrument_disposals(inst, fxbook)
        # Sales, and not a return of capital beyond the cost base, whose gain
        # (#146) is in the gross figure but is not a sale's proceeds.
        sales = [d for d in disposals if d.event == "sale"]
        if not sales:
            continue
        cost_base = _total(d.cost for d in sales)
        proceeds = _total(d.proceeds for d in sales)
        gross = _total(d.gain for d in disposals)
        discount = _total(p.discount for d in disposals for p in d.parcels)
        # The return is what came back over what was paid: indexing lowers
        # the gain taxed, not the money made (decisions.md #150).
        indexed = sum((d.indexation or ZERO for d in disposals), ZERO)
        units = sum((d.quantity for d in sales), ZERO)
        dividends_aud = _total(
            queries.in_aud(dv.cash_amount, inst.currency, fxbook.of(dv, inst.currency))
            for dv in inst.dividends
        )
        buys = [t.date for t in inst.trades if t.type in ("buy", "drp")]
        pref = prefs.get(inst.id)
        rows.append([
            inst.ticker, inst.name or "", _d(units, "0.00000001"),
            min(buys).isoformat(),          # a closed position was bought first
            max(d.date for d in sales).isoformat(),
            _d(cost_base), _d(proceeds), _d(dividends_aud), _d(gross), _d(discount),
            _d(None if gross is None or discount is None else gross - discount),
            _d(None if gross is None or dividends_aud is None
               else gross + indexed + dividends_aud),
            (pref.note if pref else "") or "",
        ])
    rows.sort(key=lambda r: r[0])
    return headers, rows


BUILDERS = {
    "transactions": transactions,
    "holdings": holdings,
    "dividends": dividends,
    "realised_cgt": realised_cgt,
    "unrealised_cgt": unrealised_cgt,
    "performance": performance,
    "closed_positions": closed_positions,
}


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #

# A spreadsheet runs a cell that starts with any of these as a formula. Notes,
# names and symbols are typed by people, so neither renderer trusts them
# (decisions.md #142).
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _inert(value):
    """Text that starts like a formula, quoted so it is shown, not run. Only
    text: a negative number is a number and keeps its sign."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([_inert(h) for h in headers])
    writer.writerows([_inert(v) for v in row] for row in rows)
    # utf-8-sig: Excel on Windows opens a plain utf-8 CSV as mojibake.
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(sheets: dict[str, Report], note: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, (headers, rows) in sheets.items():
        # Excel sheet names: 31 chars, no []:*?/\
        safe = "".join(c for c in title if c not in "[]:*?/\\")[:31]
        ws = wb.create_sheet(safe)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="left")
        for row in rows:
            ws.append(row)
        # Every text cell is text. openpyxl makes a formula of a string that
        # starts with "=", and XLSX, unlike CSV, has a type to say so, so
        # nothing is added to what the cell shows.
        for line in ws.iter_rows():
            for cell in line:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        for i, header in enumerate(headers, start=1):
            longest = max(
                [len(str(header))] + [len(str(r[i - 1])) for r in rows[:200]] or [0]
            )
            ws.column_dimensions[get_column_letter(i)].width = min(40, longest + 3)
    if note:
        ws = wb.create_sheet("About")
        for line in note.splitlines():
            ws.append([line])
        ws.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ABOUT = """Stocktake export
Generated by Stocktake, a self-hosted share portfolio tracker. All money
columns are AUD unless the header says "native"; each row is converted at its
own FX rate, or where it has none, at the nearest one stored. The FY report in
the app names any holding where that is more than a week from the day.

Capital gains: parcels are matched FIFO. The "gain after discount" columns halve
a gain on a parcel held more than twelve months. The ATO applies capital LOSSES
before that discount, the year's own and then those carried from earlier years,
which the FY report in the app does — use the FY report's net figure when
lodging, and these rows as the working.

From 1 July 2027 the discount is replaced by indexation (Treasury Laws
Amendment (Tax Reform No. 1) Act 2026). A parcel held through 30 June 2027 is
taken to be sold just before 1 July 2027 for its value at the last close and
bought again for it: its gain to then is "Gain to 30 Jun 2027", discounted as
before. A cost base from 1 July 2027 (that value, for such a parcel) is indexed
by the CPI when the parcel was held twelve months, counted from when it was
bought: "Indexed cost base". Where it changed after, the factor is the amounts'
own factors weighed together. Indexing cannot make a loss. "Cost base" stays
what the parcel cost, and "Indexation" is what indexing took off its gain, so
the gross gain is the proceeds less both. Unrealised gains are worked as a sale
today would be, with the latest CPI for this quarter's, which is not out yet.

Dividends: franking credits are only present where a statement supplied them;
blank means not recorded, not zero. The grossed-up column is cash plus franking
plus any foreign tax withheld, so it is blank on those rows too rather than
repeating the cash figure. Foreign tax withheld is income as well as cash, and
a foreign income tax offset.
"""


def build(session: Session, reports: list[str], fy: int | None = None,
          ticker: str | None = None) -> dict[str, Report]:
    # The route refuses a report it does not know before it gets here.
    return {
        TITLES[key]: BUILDERS[key](session, fy=fy, ticker=ticker)
        for key in reports
    }
