"""app/exports.py — the CSV / Excel report set.

An export is the one artefact that leaves the app: it lands in a spreadsheet, or
in an accountant's inbox, where nobody can see the code that produced it. So the
things worth freezing are shape and arithmetic.

Shape, because a row shorter than the header row does not look broken in Excel —
it silently shifts every later column up one, and a cost base lands under
"Proceeds". Arithmetic, because these numbers get lodged.

The one figure deliberately pinned twice is the CGT discount. The per-parcel
"gain after discount" column and the FY report's bottom line are *supposed* to
disagree (the ATO nets losses before the discount; the per-parcel column can't,
because it has no view of the other parcels). The About sheet says so. The gap
is asserted explicitly below so that nobody closes it thinking it's a bug.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest
from freezegun import freeze_time

import fixture_portfolio as ref
from app import exports, fyreport, queries
from factories import add_dividend, add_prices, add_trade, make_instrument

# Excel's reserved characters in a sheet name. Kept here rather than imported
# from the app so the test states the requirement independently of the code.
FORBIDDEN_SHEET_CHARS = set("[]:*?/\\")


@pytest.fixture
def portfolio(pf):
    ref.build_reference(pf)
    return pf


def cell(headers: list[str], row: list, name: str):
    """One column of one row, by header name.

    Local helper (tests must not add to the shared modules). Positional indexes
    would make these assertions unreadable and would quietly follow a column
    that moved — which is the very thing this file exists to catch.
    """
    return row[headers.index(name)]


def column(headers: list[str], rows: list[list], name: str) -> list:
    return [cell(headers, r, name) for r in rows]


# --------------------------------------------------------------------------- #
# Shape: every report, every row, rectangular
# --------------------------------------------------------------------------- #

def test_titles_and_builders_describe_the_same_report_set():
    """`build()` skips keys with no builder, so a divergence loses a report
    silently — the workbook just comes out one sheet short."""
    assert set(exports.TITLES) == set(exports.BUILDERS)


@pytest.mark.parametrize("key", sorted(exports.TITLES))
@freeze_time(ref.TODAY)
def test_every_report_builds_rectangular_rows(portfolio, key):
    headers, rows = exports.build(portfolio, [key])[exports.TITLES[key]]

    assert headers, f"{key} has no header row"
    # Every report has something to say about the reference portfolio, so an
    # empty one here means the builder stopped finding its data, not that the
    # rectangularity check below happened to have nothing to check.
    assert rows, f"{key} produced no rows for the reference portfolio"
    for i, row in enumerate(rows):
        assert len(row) == len(headers), f"{key} row {i} is ragged: {row}"


@pytest.mark.parametrize("key", sorted(exports.TITLES))
@freeze_time(ref.TODAY)
def test_every_report_builds_rectangular_rows_for_one_financial_year(portfolio, key):
    """The FY-filtered path is a separate branch in two builders and is what the
    tax-time download actually uses."""
    headers, rows = exports.build(portfolio, [key], fy=ref.FY2024)[exports.TITLES[key]]

    assert headers
    for i, row in enumerate(rows):
        assert len(row) == len(headers), f"{key} row {i} is ragged: {row}"


@pytest.mark.parametrize("key", sorted(exports.TITLES))
@freeze_time(ref.TODAY)
def test_empty_portfolio_yields_headers_and_no_rows(pf, key):
    """A brand-new account downloading a report must get an empty spreadsheet,
    not a stack trace."""
    headers, rows = exports.build(pf, [key])[exports.TITLES[key]]

    assert headers
    assert rows == []


# --------------------------------------------------------------------------- #
# CSV rendering
# --------------------------------------------------------------------------- #

def test_csv_starts_with_the_utf8_bom():
    """Without the BOM, Excel on Windows opens a utf-8 CSV as mojibake."""
    body = exports.to_csv(["Ticker", "Note"], [["ALPHA", "café"]])

    assert body[:3] == b"\xef\xbb\xbf"
    assert body.decode("utf-8-sig").startswith("Ticker,Note")


@freeze_time(ref.TODAY)
def test_csv_round_trips_through_csv_reader(portfolio):
    headers, rows = exports.transactions(portfolio)
    body = exports.to_csv(headers, rows)

    parsed = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))

    assert parsed[0] == headers
    assert len(parsed) == len(rows) + 1          # header row plus every data row
    for line in parsed[1:]:
        assert len(line) == len(headers)


def test_csv_quotes_values_containing_its_own_delimiters(pf):
    """A note is free text. A comma, a quote or a newline in it must survive the
    round trip rather than inventing extra columns."""
    inst = make_instrument(pf, "ACME", name="Acme Pty")
    add_trade(pf, inst, "2025-03-01", "buy", 10, "1.00",
              note='sold, in two "lots"\nsecond line')

    headers, rows = exports.transactions(pf)
    parsed = list(csv.reader(io.StringIO(exports.to_csv(headers, rows).decode("utf-8-sig"))))

    assert len(parsed) == 2
    assert len(parsed[1]) == len(headers)
    assert cell(headers, parsed[1], "Note") == 'sold, in two "lots"\nsecond line'


# --------------------------------------------------------------------------- #
# Workbook rendering
# --------------------------------------------------------------------------- #

def load(body: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(body))


def test_xlsx_writes_one_sheet_per_report_plus_about():
    sheets = {
        "All trades": (["A", "B"], [[1, 2], [3, 4]]),
        "Current holdings": (["X"], [[9]]),
    }

    wb = load(exports.to_xlsx(sheets, note=exports.ABOUT))

    assert wb.sheetnames == ["All trades", "Current holdings", "About"]
    assert wb["All trades"].max_row == 3      # 1 header + 2 data rows
    assert wb["All trades"].max_column == 2


def test_xlsx_omits_the_about_sheet_when_no_note_is_supplied():
    wb = load(exports.to_xlsx({"All trades": (["A"], [[1]])}))

    assert wb.sheetnames == ["All trades"]


def test_xlsx_sheet_names_are_truncated_and_stripped_of_reserved_characters():
    """Excel rejects a workbook whose sheet name is over 31 characters or holds
    any of []:*?/\\ — openpyxl would happily write one and Excel would refuse to
    open the file."""
    title = "Realised gains/losses [2024]: *draft?* \\ padding padding padding"

    wb = load(exports.to_xlsx({title: (["A"], [[1]])}))

    (name,) = wb.sheetnames
    assert not (set(name) & FORBIDDEN_SHEET_CHARS)
    # Reserved characters dropped, then cut to 31:
    #   "Realised gainslosses 2024 draft  padding padding padding"
    #    ^-- 8 + 1 + 11 + 1 + 4 + 1 + 5 = 31 chars, ending on "draft"
    assert name == "Realised gainslosses 2024 draft"
    assert len(name) == 31


@freeze_time(ref.TODAY)
def test_combined_workbook_holds_every_report(portfolio):
    sheets = exports.build(portfolio, list(exports.TITLES))

    wb = load(exports.to_xlsx(sheets, note=exports.ABOUT))

    assert wb.sheetnames == list(exports.TITLES.values()) + ["About"]
    for title, (headers, rows) in sheets.items():
        ws = wb[title]
        assert ws.max_column == len(headers), f"{title} column count"
        assert ws.max_row == len(rows) + 1, f"{title} row count"


def test_every_report_title_fits_an_excel_sheet_name():
    """The sanitiser above is a safety net; the shipped titles should never need
    it, because a truncated title is what makes two sheets collide."""
    for title in exports.TITLES.values():
        assert len(title) <= 31, title
        assert not (set(title) & FORBIDDEN_SHEET_CHARS), title


# --------------------------------------------------------------------------- #
# transactions
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_transactions_lists_every_trade_oldest_first(portfolio):
    headers, rows = exports.transactions(portfolio)

    # The reference portfolio has 9 trades: ALPHA buy + DRP, BETAX buy/buy/sell,
    # GAMMA buy, OMEGA buy, ZULU buy/sell.
    assert len(rows) == 9
    assert column(headers, rows, "Date") == sorted(column(headers, rows, "Date"))


@freeze_time(ref.TODAY)
def test_transactions_ticker_filter_returns_only_that_tickers_rows(portfolio):
    headers, rows = exports.transactions(portfolio, ticker="BETAX")

    assert column(headers, rows, "Ticker") == ["BETAX", "BETAX", "BETAX"]
    assert column(headers, rows, "Date") == ["2023-01-10", "2023-09-10", "2024-06-10"]


@freeze_time(ref.TODAY)
def test_transactions_ticker_filter_matching_nothing_returns_no_rows(portfolio):
    _headers, rows = exports.transactions(portfolio, ticker="NOVA")

    assert rows == []


@freeze_time(ref.TODAY)
def test_transactions_adds_brokerage_to_the_aud_value_of_a_buy(portfolio):
    headers, rows = exports.transactions(portfolio, ticker="BETAX")
    buy = dict(zip(headers, rows[0]))

    assert buy["Date"] == "2023-01-10"
    assert buy["Type"] == "buy"
    assert buy["Units"] == Decimal("100")
    assert buy["Price (native)"] == Decimal("5.00")
    assert buy["Brokerage (native)"] == Decimal("10.00")
    # A buy costs the parcel plus the fee: 100 x 5.00 + 10.00 = 510.00
    assert buy["Value (AUD)"] == Decimal("510.00")


@freeze_time(ref.TODAY)
def test_transactions_deducts_brokerage_from_the_aud_value_of_a_sell(portfolio):
    headers, rows = exports.transactions(portfolio, ticker="BETAX")
    sell = dict(zip(headers, rows[2]))

    assert sell["Date"] == "2024-06-10"
    assert sell["Type"] == "sell"
    # A sell nets the fee off what you receive: 150 x 8.00 - 10.00 = 1190.00
    assert sell["Value (AUD)"] == Decimal("1190.00")


@freeze_time(ref.TODAY)
def test_transactions_converts_value_at_the_trades_own_fx_rate(portfolio):
    """OMEGA is USD. The AUD column must use the rate recorded on the trade, not
    today's rate — the money changed hands at the old one."""
    headers, rows = exports.transactions(portfolio, ticker="OMEGA")
    (row,) = rows
    buy = dict(zip(headers, row))

    assert buy["Currency"] == "USD"
    assert buy["FX rate to AUD"] == Decimal("1.50")
    # (10 x 100.00 USD + 5.00 USD brokerage) x 1.50 = 1507.50 AUD
    assert buy["Value (AUD)"] == Decimal("1507.50")


@freeze_time(ref.TODAY)
def test_transactions_records_a_drp_allocation_at_its_reinvested_value(portfolio):
    """The DRP row is the units side of the distribution. It has to appear in the
    ledger — it is a real acquisition parcel — even though no cash moved."""
    headers, rows = exports.transactions(portfolio, ticker="ALPHA")
    drp = dict(zip(headers, rows[1]))

    assert drp["Type"] == "drp"
    assert drp["Units"] == Decimal("5")
    # 5 units x 10.00, no brokerage = 50.00, the distribution that bought them
    assert drp["Value (AUD)"] == Decimal("50.00")


# --------------------------------------------------------------------------- #
# realised_cgt
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_realised_cgt_lists_one_row_per_parcel_consumed(portfolio):
    headers, rows = exports.realised_cgt(portfolio, fy=ref.FY2024)

    # BETAX's sell of 150 eats two parcels; ZULU's sell of 200 eats one.
    assert column(headers, rows, "Ticker") == ["BETAX", "BETAX", "ZULU"]
    assert column(headers, rows, "Acquired") == ["2023-01-10", "2023-09-10", "2023-02-01"]
    assert set(column(headers, rows, "Financial year")) == {"FY23/24"}


@freeze_time(ref.TODAY)
def test_realised_cgt_parcel_columns_match_the_hand_computed_fifo(portfolio):
    headers, rows = exports.realised_cgt(portfolio, fy=ref.FY2024)
    old, young, zulu = (dict(zip(headers, r)) for r in rows)

    # BETAX sell 150 @ 8.00 less 10.00 brokerage: unit proceeds 1190/150 = 7.93333...
    # Parcel bought 2023-01-10 at (100 x 5.00 + 10)/100 = 5.10 a unit.
    assert old["Units"] == Decimal("100")
    assert old["Cost base (AUD)"] == Decimal("510.00")        # 100 x 5.10
    assert old["Proceeds (AUD)"] == Decimal("793.33")         # 100 x 7.93333...
    assert old["Gross gain (AUD)"] == Decimal("283.33")       # 793.33 - 510.00

    # Parcel bought 2023-09-10 at (100 x 6.00 + 10)/100 = 6.10 a unit, 50 taken.
    assert young["Units"] == Decimal("50")
    assert young["Cost base (AUD)"] == Decimal("305.00")      # 50 x 6.10
    assert young["Proceeds (AUD)"] == Decimal("396.67")       # 50 x 7.93333...
    assert young["Gross gain (AUD)"] == Decimal("91.67")      # 396.67 - 305.00

    # ZULU sell 200 @ 9.00 less 10.00: unit proceeds 1790/200 = 8.95, cost
    # (200 x 10.00 + 10)/200 = 10.05 a unit.
    assert zulu["Cost base (AUD)"] == Decimal("2010.00")      # 200 x 10.05
    assert zulu["Proceeds (AUD)"] == Decimal("1790.00")       # 200 x 8.95
    assert zulu["Gross gain (AUD)"] == Decimal("-220.00")     # a realised loss


@freeze_time(ref.TODAY)
def test_realised_cgt_days_held_drives_the_discount_flag(portfolio):
    headers, rows = exports.realised_cgt(portfolio, fy=ref.FY2024)
    old, young, zulu = (dict(zip(headers, r)) for r in rows)

    # 2023-01-10 -> 2024-06-10: 365 days to 2024-01-10, then 152 more (Jan 31 +
    # Feb 29 + Mar 31 + Apr 30 + May 31) = 517. Over twelve months, so eligible.
    assert old["Days held"] == 517
    assert old["Discount eligible"] == "yes"

    # 2023-09-10 -> 2024-06-10: 30+31+30+31+31+29+31+30+31 = 274 days. Under
    # twelve months, so the full gain is taxed.
    assert young["Days held"] == 274
    assert young["Discount eligible"] == "no"

    # ZULU was held long enough, but a loss gets no discount either way.
    assert zulu["Discount eligible"] == "yes"
    assert zulu["CGT discount (AUD)"] == Decimal("0.00")


@freeze_time(ref.TODAY)
def test_realised_cgt_gross_gains_agree_with_the_fy_report(portfolio):
    """Same disposals, same FIFO engine: the per-parcel rows must add up to the
    FY report's gross buckets, or one of the two is lying."""
    headers, rows = exports.realised_cgt(portfolio, fy=ref.FY2024)
    cgt = fyreport.fy_cgt(portfolio, ref.FY2024)

    gains = column(headers, rows, "Gross gain (AUD)")
    eligible = column(headers, rows, "Discount eligible")

    discountable = sum(g for g, e in zip(gains, eligible) if g > 0 and e == "yes")
    other = sum(g for g, e in zip(gains, eligible) if g > 0 and e == "no")
    losses = -sum(g for g in gains if g < 0)

    assert discountable == cgt.gains_discountable == ref.FY2024_GAINS_DISCOUNTABLE
    assert other == cgt.gains_other == ref.FY2024_GAINS_OTHER
    assert losses == cgt.losses == ref.FY2024_LOSSES


@freeze_time(ref.TODAY)
def test_realised_cgt_per_parcel_discount_does_not_net_losses_first(portfolio):
    """The documented caveat, pinned so it is never "fixed" by accident.

    A parcel row cannot know about the other parcels, so it halves its own gain
    whenever it was held over twelve months. The ATO applies capital losses
    BEFORE the discount, which is what the FY report does. The two totals are
    therefore meant to differ, and the About sheet tells the reader to lodge the
    FY figure. Change the per-parcel column to agree with the FY report and you
    have made every row of the working wrong.
    """
    headers, rows = exports.realised_cgt(portfolio, fy=ref.FY2024)
    cgt = fyreport.fy_cgt(portfolio, ref.FY2024)

    # Per parcel, discount applied straight to each gain:
    #   BETAX 2023-01-10  283.33 gain, eligible -> 283.33 - 141.665 = 141.665 -> 141.66
    #   BETAX 2023-09-10   91.67 gain, not eligible ->  91.67
    #   ZULU  2023-02-01 -220.00 loss, no discount -> -220.00
    assert column(headers, rows, "Gain after discount (AUD)") == [
        Decimal("141.66"), Decimal("91.67"), Decimal("-220.00"),
    ]
    per_parcel = sum(column(headers, rows, "Gain after discount (AUD)"))
    assert per_parcel == Decimal("13.33")   # 141.66 + 91.67 - 220.00

    # The FY report spends the 220.00 loss on the NON-discountable gain first
    # (91.67), then on the discountable one, leaving 283.33 - 128.33 = 155.00 to
    # be halved: 77.50.
    assert cgt.net_capital_gain == Decimal("77.50")
    assert cgt.net_capital_gain - per_parcel == Decimal("64.17")   # 77.50 - 13.33


@freeze_time(ref.TODAY)
def test_realised_cgt_fy_filter_excludes_other_years_disposals(portfolio):
    """Both disposals fall in FY23/24; no other year has any."""
    _headers, in_fy = exports.realised_cgt(portfolio, fy=ref.FY2024)
    _headers, next_fy = exports.realised_cgt(portfolio, fy=ref.FY2024 + 1)
    _headers, unfiltered = exports.realised_cgt(portfolio)

    assert len(in_fy) == 3
    assert next_fy == []
    assert len(unfiltered) == 3


# --------------------------------------------------------------------------- #
# unrealised_cgt
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_unrealised_cgt_lists_only_parcels_still_held(portfolio):
    headers, rows = exports.unrealised_cgt(portfolio)

    # ALPHA keeps both parcels (nothing sold). BETAX's FIFO sell of 150 consumed
    # the 2023-01-10 parcel whole and 50 of the 2023-09-10 one, so only the
    # younger parcel survives. GAMMA and OMEGA are untouched. ZULU is fully
    # exited and must not appear at all.
    assert [(cell(headers, r, "Ticker"), cell(headers, r, "Acquired")) for r in rows] == [
        ("ALPHA", "2021-01-04"),
        ("ALPHA", "2021-07-01"),
        ("BETAX", "2023-09-10"),
        ("GAMMA", "2024-03-01"),
        ("OMEGA", "2022-02-01"),
    ]


@freeze_time(ref.TODAY)
def test_unrealised_cgt_parcel_units_sum_to_the_current_balance(portfolio):
    headers, rows = exports.unrealised_cgt(portfolio)

    held = {}
    for row in rows:
        ticker = cell(headers, row, "Ticker")
        held[ticker] = held.get(ticker, Decimal(0)) + cell(headers, row, "Units")

    # ALPHA 100 bought + 5 DRP = 105; BETAX 200 bought - 150 sold = 50;
    # GAMMA 100; OMEGA 10; ZULU fully exited so it has no row at all.
    assert held == {
        "ALPHA": ref.UNITS["ALPHA"],
        "BETAX": ref.UNITS["BETAX"],
        "GAMMA": ref.UNITS["GAMMA"],
        "OMEGA": ref.UNITS["OMEGA"],
    }


@freeze_time(ref.TODAY)
def test_unrealised_cgt_values_each_parcel_at_todays_price(portfolio):
    headers, rows = exports.unrealised_cgt(portfolio)
    bought, drp = (dict(zip(headers, r)) for r in rows[:2])

    # The bought parcel carries the brokerage: (100 x 10.00 + 10)/100 = 10.10.
    assert bought["Cost base (AUD)"] == Decimal("1010.00")     # 100 x 10.10
    assert bought["Market value (AUD)"] == Decimal("1500.00")  # 100 x 15.00 today
    assert bought["Unrealised gain (AUD)"] == Decimal("490.00")
    # 2021-01-04 -> 2026-08-02: 1826 days to 2026-01-04 (2024 is a leap year)
    # plus 210 = 2036, so the discount would apply on a sale today.
    assert bought["Days held"] == 2036
    assert bought["Discount eligible"] == "yes"
    assert bought["Gain after discount (AUD)"] == Decimal("245.00")   # 490.00 / 2

    # The DRP parcel cost the distribution that bought it: 5 x 10.00 = 50.00.
    assert drp["Cost base (AUD)"] == Decimal("50.00")
    assert drp["Market value (AUD)"] == Decimal("75.00")       # 5 x 15.00
    assert drp["Gain after discount (AUD)"] == Decimal("12.50")       # 25.00 / 2


@freeze_time(ref.TODAY)
def test_unrealised_cgt_uses_trade_date_fx_for_cost_and_todays_fx_for_value(portfolio):
    """A foreign parcel's cost base is fixed at the rate it was bought at; only
    the market value moves with the currency."""
    headers, rows = exports.unrealised_cgt(portfolio)
    omega = dict(zip(headers, [r for r in rows if cell(headers, r, "Ticker") == "OMEGA"][0]))

    # cost: (10 x 100.00 + 5.00) USD x 1.50 = 1507.50 AUD, the 2022 rate
    assert omega["Cost base (AUD)"] == Decimal("1507.50")
    # value: 10 x 120.00 USD x 1.60 = 1920.00 AUD, today's rate
    assert omega["Market value (AUD)"] == Decimal("1920.00")
    assert omega["Unrealised gain (AUD)"] == Decimal("412.50")
    assert omega["Gain after discount (AUD)"] == Decimal("206.25")    # 412.50 / 2


@freeze_time(ref.TODAY)
def test_unrealised_cgt_gives_a_loss_no_discount(portfolio):
    headers, rows = exports.unrealised_cgt(portfolio)
    gamma = dict(zip(headers, [r for r in rows if cell(headers, r, "Ticker") == "GAMMA"][0]))

    # cost 100 x 20.10 = 2010.00, value 100 x 5.00 = 500.00
    assert gamma["Unrealised gain (AUD)"] == Decimal("-1510.00")
    # Halving a loss would understate it; the after-discount column must be the
    # loss in full even though the parcel is over twelve months old.
    assert gamma["Discount eligible"] == "yes"
    assert gamma["Gain after discount (AUD)"] == Decimal("-1510.00")


@freeze_time(ref.TODAY)
def test_unrealised_cgt_leaves_aud_value_blank_when_the_fx_rate_is_unknown(pf):
    """No USDAUD rate stored, so no AUD figure can be honest.

    Everywhere else the app refuses to guess: `queries.totals` drops such a
    holding into `excluded` rather than pretending 1 USD = 1 AUD. This report
    must not multiply by 1 and print the native number under "Market value
    (AUD)", which a spreadsheet will happily add to a real AUD column. Blank is
    the honest answer.
    """
    widget = make_instrument(pf, "WIDGET", exchange="NASDAQ", currency="USD",
                             name="Widget Inc")
    add_prices(pf, widget, [("2026-08-01", "120.00")])
    add_trade(pf, widget, "2024-01-05", "buy", 10, "100.00", fx_rate=None)

    headers, rows = exports.unrealised_cgt(pf)
    (row,) = rows

    assert queries.latest_fx(pf, "USD") is None      # the premise of the test
    assert cell(headers, row, "Market value (AUD)") == ""


# --------------------------------------------------------------------------- #
# closed_positions
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_closed_positions_lists_only_fully_exited_instruments(portfolio):
    headers, rows = exports.closed_positions(portfolio)

    # ZULU is the only position sold out; BETAX was only partly sold.
    assert column(headers, rows, "Ticker") == ["ZULU"]


@freeze_time(ref.TODAY)
def test_closed_positions_capital_gain_columns_come_from_the_fifo_engine(portfolio):
    headers, rows = exports.closed_positions(portfolio)
    zulu = dict(zip(headers, rows[0]))

    assert zulu["Units sold"] == Decimal("200")
    assert zulu["First bought"] == "2023-02-01"
    assert zulu["Last sold"] == "2024-06-20"
    # cost base 200 x (2000 + 10)/200 = 2010.00; proceeds 200 x 9.00 - 10 = 1790.00
    assert zulu["Outlay incl. brokerage (AUD)"] == Decimal("2010.00")
    assert zulu["Proceeds net of brokerage (AUD)"] == Decimal("1790.00")
    assert zulu["Gross capital gain (AUD)"] == Decimal("-220.00")
    # A loss is never discounted, so the after-discount figure is the loss itself.
    assert zulu["CGT discount (AUD)"] == Decimal("0.00")
    assert zulu["Capital gain after discount (AUD)"] == Decimal("-220.00")
    # ZULU never paid a distribution, so total return is the capital loss alone.
    assert zulu["Dividends (AUD)"] == Decimal("0.00")
    assert zulu["Total return incl. dividends (AUD)"] == Decimal("-220.00")


@freeze_time(ref.TODAY)
def test_closed_positions_gross_gain_matches_the_realised_cgt_report(portfolio):
    """Both reports are fed by `fyreport.instrument_disposals`, so a reader
    reconciling one against the other must not find two answers."""
    c_headers, c_rows = exports.closed_positions(portfolio)
    r_headers, r_rows = exports.realised_cgt(portfolio)

    zulu_gross = cell(c_headers, c_rows[0], "Gross capital gain (AUD)")
    parcels = [r for r in r_rows if cell(r_headers, r, "Ticker") == "ZULU"]

    assert zulu_gross == sum(cell(r_headers, r, "Gross gain (AUD)") for r in parcels)


# --------------------------------------------------------------------------- #
# dividends
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_dividends_puts_a_july_payment_in_the_following_financial_year(pf):
    """The AU financial year ends 30 June, so a 1 July payment is the FIRST
    income of the next year, not the last of the one just closed. Getting this
    off by a day puts a distribution on the wrong tax return.
    """
    acme = make_instrument(pf, "ACME", name="Acme Pty")
    add_trade(pf, acme, "2024-01-05", "buy", 100, "1.00")
    add_dividend(pf, acme, "2025-06-30", "10.00")   # last day of FY24/25
    add_dividend(pf, acme, "2025-07-01", "20.00")   # first day of FY25/26

    headers, fy2025 = exports.dividends(pf, fy=2025)
    _headers, fy2026 = exports.dividends(pf, fy=2026)

    assert column(headers, fy2025, "Payment date") == ["2025-06-30"]
    assert column(headers, fy2025, "Financial year") == ["FY24/25"]
    assert column(headers, fy2026, "Payment date") == ["2025-07-01"]
    assert column(headers, fy2026, "Financial year") == ["FY25/26"]


@freeze_time(ref.TODAY)
def test_dividends_unfiltered_returns_every_payment(portfolio):
    headers, rows = exports.dividends(portfolio)

    # The reference portfolio has two: ALPHA's 2021 reinvested distribution and
    # its 2022 cash one. July payments both, so both land in the NEXT FY.
    assert column(headers, rows, "Payment date") == ["2021-07-01", "2022-07-01"]
    assert column(headers, rows, "Financial year") == ["FY21/22", "FY22/23"]


@freeze_time(ref.TODAY)
def test_dividends_flags_a_reinvested_distribution(portfolio):
    """A reinvested distribution is still assessable income, but it bought units
    rather than arriving as cash — the column is what tells them apart."""
    headers, rows = exports.dividends(portfolio)

    assert column(headers, rows, "Reinvested") == ["yes", "no"]
    assert column(headers, rows, "Cash (native)") == [Decimal("50.00"), Decimal("60.00")]


@freeze_time(ref.TODAY)
def test_dividends_grosses_up_cash_plus_franking_credits(pf):
    """A fully franked dividend at the 30% company rate: the credit is
    cash x 30/70, and the grossed-up figure is what goes in the tax return."""
    acme = make_instrument(pf, "ACME", name="Acme Pty")
    add_trade(pf, acme, "2024-01-05", "buy", 100, "1.00")
    add_dividend(pf, acme, "2025-02-10", "70.00", franking_credits="30.00")

    headers, rows = exports.dividends(pf)
    div = dict(zip(headers, rows[0]))

    assert div["Cash (AUD)"] == Decimal("70.00")
    assert div["Franking credits (AUD)"] == Decimal("30.00")
    assert div["Grossed up (AUD)"] == Decimal("100.00")   # 70.00 + 30.00


@freeze_time(ref.TODAY)
def test_dividends_converts_foreign_cash_at_the_payments_own_rate(pf):
    nova = make_instrument(pf, "NOVA", exchange="NASDAQ", currency="USD", name="Nova Inc")
    add_trade(pf, nova, "2024-01-05", "buy", 100, "1.00", fx_rate="1.40")
    add_dividend(pf, nova, "2025-02-10", "100.00", fx_rate="1.50")

    headers, rows = exports.dividends(pf)
    div = dict(zip(headers, rows[0]))

    assert div["Cash (native)"] == Decimal("100.00")
    assert div["Currency"] == "USD"
    assert div["Cash (AUD)"] == Decimal("150.00")    # 100.00 USD x 1.50


@freeze_time(ref.TODAY)
def test_dividends_leaves_franking_blank_when_it_was_never_recorded(pf):
    """`ABOUT` promises: "blank means not recorded, not zero".

    The distinction is the whole reason `fyreport.IncomeRow` carries a
    `franking_missing` count. An export that writes 0.00 for both cases turns
    "we don't know yet" into "there was none", and the reader has no way to see
    which credits are still missing from their return.
    """
    acme = make_instrument(pf, "ACME", name="Acme Pty")
    add_trade(pf, acme, "2024-01-05", "buy", 100, "1.00")
    add_dividend(pf, acme, "2025-02-10", "70.00")           # franking not known
    add_dividend(pf, acme, "2025-03-10", "80.00", franking_credits="0")  # unfranked

    headers, rows = exports.dividends(pf)

    assert column(headers, rows, "Franking credits (AUD)") == ["", Decimal("0.00")]
    # The grossed-up figure is cash + franking, so it is unknowable for the
    # first row and exactly the cash for the second. Writing 70.00 against an
    # unrecorded credit would be the same false claim one column across.
    assert column(headers, rows, "Grossed up (AUD)") == ["", Decimal("80.00")]


# --------------------------------------------------------------------------- #
# performance
# --------------------------------------------------------------------------- #

@freeze_time(ref.TODAY)
def test_performance_has_one_row_per_financial_year(portfolio):
    """The first trade is 2021-01-04 (FY20/21) and today is 2026-08-02, which is
    already in FY26/27 — seven financial years touched, oldest first."""
    headers, rows = exports.performance(portfolio)

    assert column(headers, rows, "Financial year") == [
        "FY20/21", "FY21/22", "FY22/23", "FY23/24", "FY24/25", "FY25/26", "FY26/27",
    ]
    assert len(rows) == len(queries.annual_series(portfolio)["labels"])


@freeze_time(ref.TODAY)
def test_performance_first_year_matches_the_hand_computed_series(portfolio):
    headers, rows = exports.performance(portfolio)
    first = dict(zip(headers, rows[0]))

    # FY20/21 holds one trade: ALPHA 100 @ 10.00 + 10.00 brokerage on 2021-01-04.
    # The last price point in that FY is 2021-06-01, where ALPHA closes at 10.00.
    assert first["Invested in year (AUD)"] == Decimal("1010.00")
    assert first["Invested cumulative (AUD)"] == Decimal("1010.00")
    assert first["Value at year end (AUD)"] == Decimal("1000.00")   # 100 x 10.00
    assert first["Net gain cumulative (AUD)"] == Decimal("-10.00")  # the brokerage
    # -10.00 / 1010.00 = -0.00990099..., rounded to 4dp then shown as a percent
    assert first["Gain %"] == Decimal("-0.99")
    assert first["Gain in year (AUD)"] == Decimal("-10.00")
    assert first["Year gain %"] == Decimal("-0.99")


@freeze_time(ref.TODAY)
def test_performance_final_year_matches_the_hand_computed_series(portfolio):
    headers, rows = exports.performance(portfolio)
    last = dict(zip(headers, rows[-1]))

    assert last["Financial year"] == "FY26/27"
    # Nothing was bought after 2024-03-01, so the current FY added no capital.
    assert last["Invested in year (AUD)"] == Decimal("0.00")
    assert last["Invested cumulative (AUD)"] == ref.SERIES_INVESTED     # 7657.50
    assert last["Value at year end (AUD)"] == ref.SERIES_VALUE          # 4445.00
    assert last["Net gain cumulative (AUD)"] == ref.SERIES_GAIN         # -172.50
    # -172.50 / 7657.50 = -0.0225269..., rounded to 4dp = -0.0225
    assert last["Gain %"] == Decimal("-2.25")


@freeze_time(ref.TODAY)
def test_performance_cumulative_columns_only_move_forward(portfolio):
    """Cumulative invested is a running total of buys, so it can never fall —
    a decreasing value would mean sale proceeds had been netted off it, which is
    the mistake that made the all-time return disagree with the yearly chart."""
    headers, rows = exports.performance(portfolio)
    cumulative = column(headers, rows, "Invested cumulative (AUD)")

    assert cumulative == sorted(cumulative)
